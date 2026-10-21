# -*- coding: utf-8 -*-
"""临时验证：AI 审核力度增强 + AI 问题/提示与本地规则问题合并进报告。"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from checkers.base import Issue, FileResult, STATUS_ISSUE, STATUS_PASS  # noqa: E402
from checkers.scanner import summarize  # noqa: E402
from checkers import ai_checker  # noqa: E402
from report.report_builder import build_report  # noqa: E402

ok = fail = 0


def chk(name, cond, extra=""):
    global ok, fail
    if cond:
        ok += 1
        print(f"  [OK] {name}")
    else:
        fail += 1
        print(f"  [FAIL] {name} {extra}")


# ---------------------------------------------------------------- 1. 引擎增强
print("\n== 1. AI 核验引擎增强 ==")
blocks = [(f"第{i}段", "内容" * 400) for i in range(1, 41)]   # 40 段 × 800 字
chunks = ai_checker._chunk_blocks(blocks, 3000, 10)
locs = [c[0] for c in chunks]
chk("分段数量受 max_requests 限制", len(chunks) == 10, f"实际 {len(chunks)}")
chk("长文档覆盖首段", locs[0] == "第1段", locs[:3])
chk("长文档覆盖末段（不再只检前 N 段）", locs[-1] == "第40段", locs[-3:])
chk("全文均匀抽样", len(set(locs)) == 10, locs)

src = "本报告评估基准日为2022年10月31日，账面价值合计为300万元。"
chk("严格取证：真实引用通过", ai_checker._quote_grounded("评估基准日为2022年10月31日", src))
chk("严格取证：编造引用被拦截", not ai_checker._quote_grounded("评估基准日为2023年12月31日", src))
chk("严格取证：短引用需整段命中", not ai_checker._quote_grounded("基准日2021", src))

payload = (
    '[{"quote":"账面价值合计为300万元","category":"表格计算",'
    '"issue":"合计与分项不符","detail":"分项80+120=200，合计写为300",'
    '"suggestion":"将合计改为200万元","severity":"high","confidence":0.9},'
    '{"quote":"评估基准日","category":"数据错误","issue":"日期错误",'
    '"detail":"","suggestion":"核实日期","severity":"low","confidence":0.2}]'
)
items = ai_checker._parse_issues(payload, src, 0.6)
chk("置信度门槛过滤低把握条目", len(items) == 1, [i["issue"] for i in items])
chk("保留条目带类别与举证", items and items[0]["category"] == "表格计算"
    and "80+120" in items[0]["detail"])

rec = ai_checker.build_prompt_record({"enabled": True, "mode": "local", "model": "qwen3:8b"})
chk("提示词记录含主提示词", "逐句、逐单元格通读" in rec["system_prompt"])
chk("提示词记录含 9 类维度", "交叉一致性" in rec["system_prompt"] and "规范符合性" in rec["system_prompt"])
chk("提示词记录含深度复核提示", "深度复核" in (rec["review_prompt"] or ""))
chk("提示词记录含置信度条款", "confidence" in rec["system_prompt"])
chk("深度复核默认开启", rec["deep_review"] is True)

# ------------------------------------------------------- 2. 报告合并汇总
print("\n== 2. AI 问题 / 提示与本地规则问题合并进报告 ==")
res1 = FileResult("评估报告.docx", "D:/x/评估报告.docx", "Word", 12345, STATUS_ISSUE, issues=[
    Issue(rule_key="full_width_space", rule_title="多余全角空格", severity="low",
          location="第 3 段", detail="存在多余全角空格", snippet="评估　报告",
          suggestion="删除多余空格", source="builtin"),
    Issue(rule_key="ai_verify", rule_title="AI 智能核验", severity="high",
          location="第 12 段", detail="[表格计算] 合计与分项不符：80+120=200，文中写为300",
          snippet="合计为300万元", suggestion="将合计改为200万元", source="ai", tag="表格计算"),
    Issue(rule_key="ai_verify", rule_title="AI 智能核验", severity="medium",
          location="第 20 段", detail="[用词不当] 术语误用：应为『评估基准日』",
          snippet="评估基准日期", suggestion="统一为『评估基准日』", source="ai", tag="用词不当"),
])
res2 = FileResult("明细表.xlsx", "D:/x/明细表.xlsx", "Excel", 8888, STATUS_ISSUE, issues=[
    Issue(rule_key="num_format", rule_title="数字格式不一致", severity="medium",
          location="Sheet1!C5", detail="数字格式与同列不一致", snippet="1,234",
          suggestion="统一为两位小数", source="builtin"),
])
res2.ai_status = "partial"
res2.ai_note = "部分分段失败：AI 响应超时（已命中 0 条）"
res3 = FileResult("说明.pdf", "D:/x/说明.pdf", "PDF", 555, STATUS_PASS)
res3.ai_status = "ok"
res3.ai_note = "AI 核验完成：命中 0 条问题"

results = [res1, res2, res3]
summary = summarize(results)
chk("汇总已包含 AI 问题（total_issues=4）", summary["total_issues"] == 4, summary["total_issues"])
chk("AI 问题进入严重级别统计", summary["severity"]["high"] == 1, summary["severity"])

out = os.path.join(tempfile.gettempdir(), "_t_ai_merge_report.docx")
build_report(results, summary, out, rule_summary={"word": 10, "excel": 8, "textnorm": 6, "pdf": 4},
             operator="张三", org="测试单位", ai_review=rec)

from docx import Document  # noqa: E402
d = Document(out)
texts = [p.text for p in d.paragraphs]
all_text = "\n".join(texts)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            all_text += "\n" + c.text

chk("报告含「六、AI 智能核验专项汇总」", "六、AI 智能核验专项汇总" in all_text)
chk("报告含「七、整改建议」", "七、整改建议" in all_text)
chk("报告含 2.3 问题来源分布", "2.3 问题来源分布" in all_text)
chk("检测方式标注 AI 参与", "本地规则校验 + AI 智能核验" in all_text)
chk("AI 提示词写入报告", "逐句、逐单元格通读" in all_text)
chk("深度复核提示词写入报告", "【AI 深度复核提示词】" in all_text)
chk("模型信息写入报告", "qwen3:8b" in all_text)
chk("AI 分类统计表存在", "AI 发现问题分类统计" in all_text and "表格计算" in all_text)
chk("AI 逐条清单存在", "AI 问题逐条清单" in all_text and "合计与分项不符" in all_text)
chk("AI 提示与建议汇总存在", "AI 提示与建议汇总" in all_text and "将合计改为200万元" in all_text)
chk("未完成 AI 核验文件入表", "未完成 AI 核验的文件" in all_text and "明细表.xlsx" in all_text)
chk("明细表含「来源」列", "来源" in all_text and "本地规则" in all_text)
chk("问题类型分布按 AI 类别展开", "AI 智能核验 · 表格计算" in all_text)
chk("整改建议并入 AI 提示", "AI 智能核验共给出" in all_text)
chk("文件明细含 AI 核验结论", "AI 智能核验：部分完成" in all_text)

print(f"\n结果：{ok} 通过 / {fail} 失败")
sys.exit(1 if fail else 0)
