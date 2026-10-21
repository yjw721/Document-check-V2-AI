# -*- coding: utf-8 -*-
"""验证：专项技能（SKILL）写入 AI 核验报告（第六章 6.1 配置与提示词记录）。"""
import os
import sys
import tempfile

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))

from checkers.base import Issue, FileResult, STATUS_ISSUE  # noqa: E402
from checkers import ai_checker  # noqa: E402
from report.report_builder import build_report  # noqa: E402
from docx import Document  # noqa: E402

ok = fail = 0


def chk(name, cond, extra=""):
    global ok, fail
    if cond:
        ok += 1
        print(f"  [OK] {name}")
    else:
        fail += 1
        print(f"  [FAIL] {name} {extra}")


ai_cfg = {
    "enabled": True,
    "mode": "local",
    "model": "qwen3:8b",
    "active_skills": ["financial_reconciliation", "terminology_consistency", "sk1"],
    "custom_skills": [{
        "id": "sk1", "name": "合同风险点", "prompt": "核查违约责任与交付节点是否明确。",
        "file_types": "doc", "default_severity": "high",
    }],
}
rec = ai_checker.build_prompt_record(ai_cfg)
chk("记录含 active_skills", rec.get("skill_count") == 3, rec.get("skill_count"))
chk("记录含 skill_prompts", bool(rec.get("skill_prompts")), rec.get("skill_prompts"))

res1 = FileResult("评估报告.docx", "D:/x/评估报告.docx", "Word", 12345, STATUS_ISSUE, issues=[
    Issue(rule_key="ai_verify", rule_title="AI 智能核验", severity="high",
          location="第 12 段", detail="[表格计算] 合计与分项不符：80+120=200，文中写为300",
          snippet="合计为300万元", suggestion="将合计改为200万元", source="ai", tag="表格计算"),
])
summary = {"total_files": 1, "total_issues": 1, "severity": {"high": 1, "medium": 0, "low": 0},
           "by_rule": {}, "by_source": {}}

out = os.path.join(tempfile.gettempdir(), "_t_skill_report.docx")
build_report([res1], summary, out, rule_summary={"word": 1},
             operator="测试", org="测试单位", ai_review=rec)

d = Document(out)
all_text = "\n".join(p.text for p in d.paragraphs)
for t in d.tables:
    for row in t.rows:
        for c in row.cells:
            all_text += "\n" + c.text

chk("6.1 含「启用专项技能」行", "启用专项技能" in all_text)
chk("报告含内置技能名「财务数据勾稽」", "财务数据勾稽" in all_text)
chk("报告含自定义技能名「合同风险点」", "合同风险点" in all_text)
chk("报告含专项技能提示词区块", "【启用的专项技能提示词】" in all_text)
chk("报告含自定义技能提示词正文", "核查违约责任与交付节点是否明确" in all_text)

print(f"\n结果：{ok} 通过 / {fail} 失败")
sys.exit(1 if fail else 0)
