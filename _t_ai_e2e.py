# -*- coding: utf-8 -*-
"""端到端：上传样例 → 本地 AI 核验 → 导出报告 → 校验 AI 问题/提示已合并进报告。"""
import json
import os
import sys
import time
import urllib.request
import uuid

BASE = "http://127.0.0.1:8501"
TMP = os.path.join(os.path.dirname(os.path.abspath(__file__)), "tests", "_tmp_e2e")
os.makedirs(TMP, exist_ok=True)

from docx import Document  # noqa: E402

# 1. 造一个含「全角空格」+「合计算错」的样例
sample = os.path.join(TMP, "AI合并汇总样例.docx")
d = Document()
d.add_paragraph("资产评估报告（样例）")
d.add_paragraph("本次评估基准日为　2022年10月31日，评估范围为确山江祥风力发电有限公司的全部资产。")
d.add_paragraph("账面价值合计为300万元，其中：房屋建筑物80万元，机器设备120万元。")
d.add_paragraph("评估结论：委估资产在评估基准日的市场价值为叁佰万元整。")
d.save(sample)


def post_file(url, path, extra=b""):
    boundary = "----wb" + uuid.uuid4().hex
    name = os.path.basename(path)
    with open(path, "rb") as fp:
        content = fp.read()
    body = b""
    body += f"--{boundary}\r\n".encode()
    body += (f'Content-Disposition: form-data; name="files"; filename="{name}"\r\n'
             f"Content-Type: application/octet-stream\r\n\r\n").encode()
    body += content + b"\r\n"
    body += f"--{boundary}--\r\n".encode()
    req = urllib.request.Request(url, data=body, method="POST")
    req.add_header("Content-Type", f"multipart/form-data; boundary={boundary}")
    with urllib.request.urlopen(req, timeout=120) as r:
        return json.loads(r.read().decode("utf-8"))


def get(url, timeout=30):
    with urllib.request.urlopen(url, timeout=timeout) as r:
        return json.loads(r.read().decode("utf-8"))


print("1) 上传并启动核验 …")
up = post_file(f"{BASE}/api/upload", sample)
tid = up.get("task_id")
print("   task_id =", tid)
assert tid, up

deadline = time.time() + 900
last_stage = ""
while time.time() < deadline:
    t = get(f"{BASE}/api/task/{tid}")
    if t.get("stageText") and t["stageText"] != last_stage:
        last_stage = t["stageText"]
        print(f"   [{t.get('progress')}%] {t.get('stage')} - {last_stage}")
    if t.get("status") in ("done", "error", "cancelled"):
        print("   最终状态:", t.get("status"), t.get("error", ""))
        break
    time.sleep(2)

iss = get(f"{BASE}/api/issues")
items = iss.get("issues", [])
ai_items = [i for i in items if i.get("source") == "ai"]
rule_items = [i for i in items if i.get("source") != "ai"]
print(f"2) 问题总数 {len(items)}：本地规则 {len(rule_items)} / AI 核验 {len(ai_items)}")
for i in ai_items[:10]:
    print("   AI:", i.get("tag"), "|", i.get("severity"), "|", (i.get("detail") or "")[:90])

print("3) 导出报告 …")
body = json.dumps({"operator": "端到端测试", "org": "测试单位",
                   "include_cover": True}).encode("utf-8")
req = urllib.request.Request(f"{BASE}/api/report", data=body, method="POST")
req.add_header("Content-Type", "application/json")
with urllib.request.urlopen(req, timeout=180) as r:
    raw = r.read()
out = os.path.join(TMP, "AI合并汇总报告.docx")
with open(out, "wb") as fp:
    fp.write(raw)
print("   报告已保存：", out, len(raw), "字节")

doc = Document(out)
all_text = "\n".join(p.text for p in doc.paragraphs)
for t in doc.tables:
    for row in t.rows:
        for c in row.cells:
            all_text += "\n" + c.text

checks = [
    ("六、AI 智能核验专项汇总", "六、AI 智能核验专项汇总" in all_text),
    ("七、整改建议", "七、整改建议" in all_text),
    ("2.3 问题来源分布", "2.3 问题来源分布" in all_text),
    ("检测方式含 AI", "本地规则校验 + AI 智能核验" in all_text),
    ("AI 主提示词归档", "逐句、逐单元格通读" in all_text),
    ("AI 复核提示词归档", "【AI 深度复核提示词】" in all_text),
    ("模型名归档", "qwen3:8b" in all_text),
    ("来源列/本地规则", "本地规则" in all_text),
]
bad = 0
for name, cond in checks:
    print(f"   [{'OK' if cond else 'FAIL'}] {name}")
    if not cond:
        bad += 1
if ai_items:
    hit = any((i.get("suggestion") or "")[:12] in all_text for i in ai_items)
    print(f"   [{'OK' if hit else 'FAIL'}] AI 提示进入报告 6.4 汇总")
    if not hit:
        bad += 1
print("\n结果：", "全部通过" if bad == 0 else f"{bad} 项未通过")
sys.exit(1 if bad else 0)
