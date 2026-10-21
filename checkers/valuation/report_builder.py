# -*- coding: utf-8 -*-
"""
资产评估报告智能核查 —— Word 终稿核查报告生成（纯本地 python-docx）
=========================================================================
报告结构（对应《AI核查报告需求点提取》四阶段闭环）：
    封面 → 一、核查基本信息 → 二、四阶段执行概览
    → 三、整体统计汇总 → 四、问题明细清单（按节点分组）
    → 五、去重与 AI 复核说明 → 六、整改建议与结论 → 落款

排版特性（与 report/report_builder.py 保持一致）：
    A4 页面、中文字体（微软雅黑 + 宋体）、带边框表格、斑马纹表头、
    标题分隔线、页脚页码、跨页重复表头。

保密说明：仅本机生成 docx 文件，不联网、不上传、不写入任何外部位置。
"""

from __future__ import annotations

import os
from datetime import datetime
from typing import Any, Dict, List, Optional

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

# ---------------------------------------------------------------------------
# 样式常量（与项目现有报告风格一致）
# ---------------------------------------------------------------------------
FONT_TITLE = "微软雅黑"
FONT_BODY = "宋体"
COLOR_PRIMARY = RGBColor(0x1F, 0x4E, 0x79)     # 沉稳科技蓝
COLOR_TEXT = RGBColor(0x33, 0x33, 0x33)
COLOR_MUTED = RGBColor(0x80, 0x80, 0x80)
COLOR_HIGH = RGBColor(0xC0, 0x39, 0x2B)
COLOR_MEDIUM = RGBColor(0xB9, 0x7A, 0x0A)
COLOR_LOW = RGBColor(0x55, 0x7A, 0x95)

SHADE_HEADER = "DCE6F1"    # 表头底色（浅蓝）
SHADE_SUBTLE = "F2F5F9"    # 次级底色
SHADE_CHECKED = "EAF3EA"   # 已核查 / 复核修正行浅绿

SEVERITY_LABEL = {"high": "严重", "medium": "一般", "low": "轻微"}
SEVERITY_COLOR = {"high": COLOR_HIGH, "medium": COLOR_MEDIUM, "low": COLOR_LOW}
SEVERITY_RANK = {"high": 0, "medium": 1, "low": 2}


# ---------------------------------------------------------------------------
# 底层排版辅助（与 report_builder 同源，保持报告风格统一）
# ---------------------------------------------------------------------------
def _set_run_font(run: Any, name: str = FONT_BODY, size: float = 10.5,
                  bold: bool = False, color: Optional[RGBColor] = None) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:eastAsia"), name)
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)


def _para(doc_or_cell: Any, text: str = "", size: float = 10.5, bold: bool = False,
          align: Optional[int] = None, font: str = FONT_BODY,
          color: Optional[RGBColor] = None, space_after: float = 6,
          space_before: float = 0, line_spacing: Optional[float] = None,
          first_line_indent: Optional[float] = None) -> Any:
    p = doc_or_cell.add_paragraph()
    if align is not None:
        p.alignment = align
    pf = p.paragraph_format
    pf.space_after = Pt(space_after)
    pf.space_before = Pt(space_before)
    if line_spacing:
        pf.line_spacing = line_spacing
    if first_line_indent:
        pf.first_line_indent = Pt(first_line_indent)
    if text:
        run = p.add_run(text)
        _set_run_font(run, font, size, bold, color)
    return p


def _shade_cell(cell: Any, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def _set_cell_text(cell: Any, text: str, size: float = 9.5, bold: bool = False,
                   align: Optional[int] = None, color: Optional[RGBColor] = None,
                   font: str = FONT_BODY) -> None:
    cell.text = ""
    p = cell.paragraphs[0]
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.space_before = Pt(2)
    run = p.add_run("" if text is None else str(text))
    _set_run_font(run, font, size, bold, color)


def _set_repeat_header(row: Any) -> None:
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement("w:tblHeader")
    tblHeader.set(qn("w:val"), "true")
    trPr.append(tblHeader)


def _set_col_widths(table: Any, widths_cm: List[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_cm):
            if idx < len(row.cells):
                row.cells[idx].width = Cm(width)


def _add_page_number_footer(section: Any) -> None:
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def _field(instr: str) -> None:
        run = p.add_run()
        fld_begin = OxmlElement("w:fldChar")
        fld_begin.set(qn("w:fldCharType"), "begin")
        instr_el = OxmlElement("w:instrText")
        instr_el.set(qn("xml:space"), "preserve")
        instr_el.text = instr
        fld_end = OxmlElement("w:fldChar")
        fld_end.set(qn("w:fldCharType"), "end")
        run._r.append(fld_begin)
        run._r.append(instr_el)
        run._r.append(fld_end)
        _set_run_font(run, FONT_BODY, 9, color=COLOR_MUTED)

    r1 = p.add_run("第 ")
    _set_run_font(r1, FONT_BODY, 9, color=COLOR_MUTED)
    _field("PAGE")
    r2 = p.add_run(" 页  共 ")
    _set_run_font(r2, FONT_BODY, 9, color=COLOR_MUTED)
    _field("NUMPAGES")
    r3 = p.add_run(" 页")
    _set_run_font(r3, FONT_BODY, 9, color=COLOR_MUTED)


def _h1(doc: Any, text: str) -> None:
    p = _para(doc, text, size=15, bold=True, font=FONT_TITLE,
              color=COLOR_PRIMARY, space_before=14, space_after=8)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1F4E79")
    pBdr.append(bottom)
    pPr.append(pBdr)


def _h2(doc: Any, text: str) -> None:
    _para(doc, text, size=12, bold=True, font=FONT_TITLE,
          color=COLOR_TEXT, space_before=10, space_after=5)


def _kv_table(doc: Any, rows: List[tuple], key_width: float = 4.5,
              val_width: float = 11.5) -> Any:
    table = doc.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for key, val in rows:
        cells = table.add_row().cells
        _set_cell_text(cells[0], key, size=10, bold=True)
        _shade_cell(cells[0], SHADE_SUBTLE)
        _set_cell_text(cells[1], val, size=10)
    _set_col_widths(table, [key_width, val_width])
    return table


def _table_header(table: Any, headers: List[str], widths: List[float]) -> None:
    for idx, head in enumerate(headers):
        _set_cell_text(table.rows[0].cells[idx], head, size=9.5, bold=True,
                       align=WD_ALIGN_PARAGRAPH.CENTER, font=FONT_TITLE)
        _shade_cell(table.rows[0].cells[idx], SHADE_HEADER)
    _set_repeat_header(table.rows[0])
    _set_col_widths(table, widths)


# ---------------------------------------------------------------------------
# 报告主体
# ---------------------------------------------------------------------------
class ValuationReportBuilder:
    """资产评估智能核查 · Word 终稿报告生成器。"""

    def __init__(self, result: Dict[str, Any], checklist: Optional[Dict[str, Any]] = None,
                 operator: str = "", org: str = "", include_cover: bool = True) -> None:
        self.result = result or {}
        self.checklist = checklist or {}
        self.operator = operator or "—"
        self.org = org or "—"
        self.include_cover = include_cover
        self.now = datetime.now()
        self.summary = self.result.get("summary", {}) or {}
        self.findings: List[Dict[str, Any]] = list(self.result.get("findings", []) or [])

    # ------------------------------------------------------------------
    def build(self, out_path: str) -> str:
        doc = Document()
        self._setup_page(doc)
        if self.include_cover:
            self._cover(doc)
        self._basic_info(doc)
        self._stage_overview(doc)
        self._overall_stats(doc)
        self._issue_details(doc)
        self._dedup_review(doc)
        self._suggestions(doc)
        self._footer_note(doc)

        out_dir = os.path.dirname(os.path.abspath(out_path))
        if out_dir and not os.path.isdir(out_dir):
            os.makedirs(out_dir, exist_ok=True)
        doc.save(out_path)
        return out_path

    # ------------------------------------------------------------------
    def _setup_page(self, doc: Any) -> None:
        section = doc.sections[0]
        section.page_width = Cm(21.0)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(3.0)
        section.right_margin = Cm(3.0)
        _add_page_number_footer(section)
        style = doc.styles["Normal"]
        style.font.name = FONT_BODY
        style.font.size = Pt(10.5)
        rPr = style.element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = OxmlElement("w:rFonts")
            rPr.append(rFonts)
        rFonts.set(qn("w:eastAsia"), FONT_BODY)

    # ------------------------------------------------------------------
    def _cover(self, doc: Any) -> None:
        _para(doc, "", space_after=54)
        _para(doc, "资产评估报告智能核查报告", size=28, bold=True, font=FONT_TITLE,
              align=WD_ALIGN_PARAGRAPH.CENTER, color=COLOR_PRIMARY, space_after=10)
        _para(doc, "Asset Valuation Report Intelligent Verification Report", size=11,
              font=FONT_TITLE, align=WD_ALIGN_PARAGRAPH.CENTER, color=COLOR_MUTED,
              space_after=8)
        _para(doc, "— 本地离线核查 · 四阶段闭环 · 保密模式 —", size=11, font=FONT_TITLE,
              align=WD_ALIGN_PARAGRAPH.CENTER, color=COLOR_MUTED, space_after=44)

        status = self.result.get("status", "done")
        total = int(self.summary.get("total", 0))
        high = int(self.summary.get("high", 0))

        table = doc.add_table(rows=0, cols=2)
        table.style = "Table Grid"
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        cover_rows = [
            ("核查节点", f"{int(self.summary.get('node_done', 0))} / {int(self.summary.get('node_total', 0))} 个"),
            ("核查事项", f"{int(self.summary.get('item_done', 0))} / {int(self.summary.get('item_total', 0))} 项"),
            ("发现问题", f"{total} 条（严重 {high} 条）"),
            ("核查状态", self._status_label(status)),
            ("检测时间", self.now.strftime("%Y 年 %m 月 %d 日 %H:%M")),
            ("检测人", self.operator),
            ("所属单位", self.org),
        ]
        for key, val in cover_rows:
            cells = table.add_row().cells
            _set_cell_text(cells[0], key, size=11, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            _shade_cell(cells[0], SHADE_HEADER)
            _set_cell_text(cells[1], val, size=11, align=WD_ALIGN_PARAGRAPH.CENTER)
        _set_col_widths(table, [5.5, 7.5])

        _para(doc, "", space_after=36)
        _para(doc, "本报告由本地离线核查工具自动生成，全过程无网络连接、无数据外发，",
              size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=COLOR_MUTED, space_after=2)
        _para(doc, "所有文档内容仅在本机内存与本地磁盘处理，报告仅供内部核查使用。",
              size=9, align=WD_ALIGN_PARAGRAPH.CENTER, color=COLOR_MUTED, space_after=0)
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    # ------------------------------------------------------------------
    def _basic_info(self, doc: Any) -> None:
        _h1(doc, "一、核查基本信息")
        slots_meta = self.checklist.get("slots") or {}
        slot_items = []
        for slot in ("report", "statement", "income", "cost", "market"):
            meta = slots_meta.get(slot) or {}
            name = meta.get("name") or slot
            slot_items.append(f"{name}")
        # 输入附件文件名（来自预处理结果）
        pre_items = {it.get("slot"): it for it in (self.result.get("preprocess", {}).get("items") or [])}
        files_line = "、".join(
            (pre_items.get(s, {}).get("file") or (slots_meta.get(s, {}).get("name") or s))
            for s in ("report", "statement", "income", "cost", "market")
        ) or "—"

        _kv_table(doc, [
            ("核查工具", "资产评估报告智能核查（本地离线版）"),
            ("运行模式", "纯本地离线运行，零联网、零数据上传、零云端接口"),
            ("检测时间", self.now.strftime("%Y-%m-%d %H:%M:%S")),
            ("输入附件", files_line),
            ("核查节点", f"{int(self.summary.get('node_total', 0))} 个（9 大节点：报告/说明/收益法/成本法/市场法 5 个单查 + 报告↔说明、说明↔三法 4 个联查）"),
            ("核查事项", f"{int(self.summary.get('item_total', 0))} 项（96 项核查要点）"),
            ("检测方式", "本地确定性规则引擎 + AI 语义核查（可选，本地模型）+ AI 全局复核"),
            ("原文档处理", "只读核查，不对原始文档做任何写入或修改；问题命中即阻断（#REF! 等硬性错误）"),
            ("检测人", self.operator),
            ("所属单位", self.org),
        ])
        _para(doc, "", space_after=4)

    # ------------------------------------------------------------------
    def _stage_overview(self, doc: Any) -> None:
        _h1(doc, "二、四阶段执行概览")
        stages = self.checklist.get("stages") or [
            {"no": "Ⅰ", "name": "文件上传与预处理"},
            {"no": "Ⅱ", "name": "文件智能切分"},
            {"no": "Ⅲ", "name": "9 大核查节点"},
            {"no": "Ⅳ", "name": "汇总与复核"},
        ]
        for st in stages:
            _h2(doc, f"阶段 {st.get('no', '')} · {st.get('name', '')}")

        # S1 预处理
        pre = self.result.get("preprocess", {}) or {}
        if pre.get("items"):
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            _table_header(table, ["附件槽位", "文件", "预处理结果", "深度清理"],
                          [2.2, 4.3, 5.5, 3.5])
            for it in pre["items"]:
                clean = it.get("clean") or {}
                clean_txt = "无"
                if clean:
                    parts = []
                    if clean.get("hidden_sheets"):
                        parts.append(f"隐藏表 {len(clean['hidden_sheets'])}")
                    if clean.get("defined_names"):
                        parts.append(f"名称 {clean['defined_names']}")
                    if clean.get("external_links"):
                        parts.append(f"外链 {clean['external_links']}")
                    clean_txt = "；".join(parts) or "无"
                cells = table.add_row().cells
                _set_cell_text(cells[0], str(it.get("slot") or ""), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                _set_cell_text(cells[1], str(it.get("file") or ""), size=9)
                _set_cell_text(cells[2], str(it.get("message") or ""), size=9)
                _set_cell_text(cells[3], clean_txt, size=9)
        else:
            _para(doc, pre.get("message") or "（无预处理明细）", size=10, line_spacing=1.5)
        if pre.get("blocked"):
            _para(doc, "⚠ " + (pre.get("blocked_message") or "预处理未通过，核查已终止。"),
                  size=10, color=COLOR_HIGH, space_before=4)
        _para(doc, "", space_after=4)

        # S2 切分
        split = self.result.get("split", {}) or {}
        blocks = split.get("blocks") or []
        if blocks:
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            _table_header(table, ["数据块", "名称", "类型", "规模"], [1.8, 4.2, 2.2, 7.3])
            for b in blocks:
                if b.get("kind") == "excel":
                    own = (b.get("stats") or {}).get("own_sheets") or []
                    gen = b.get("general_sheets") or []
                    scale = f"自有表 {len(own)} + 通用表 {len(gen)}"
                else:
                    scale = f"{b.get('chars', 0)} 字"
                cells = table.add_row().cells
                _set_cell_text(cells[0], str(b.get("id") or ""), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                _set_cell_text(cells[1], str(b.get("name") or ""), size=9)
                _set_cell_text(cells[2], "Excel" if b.get("kind") == "excel" else "文档", size=9,
                               align=WD_ALIGN_PARAGRAPH.CENTER)
                _set_cell_text(cells[3], scale, size=9)
        _para(doc, "", space_after=4)

        # S3 节点
        nodes = self.result.get("nodes") or []
        if nodes:
            table = doc.add_table(rows=1, cols=5)
            table.style = "Table Grid"
            _table_header(table, ["节点", "名称", "状态", "事项完成", "命中数"],
                          [1.6, 5.5, 2.2, 3.0, 3.2])
            for n in nodes:
                cells = table.add_row().cells
                _set_cell_text(cells[0], str(n.get("node_no") or ""), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                _set_cell_text(cells[1], str(n.get("node_name") or ""), size=9)
                _set_cell_text(cells[2], self._status_label(n.get("status")), size=9,
                               align=WD_ALIGN_PARAGRAPH.CENTER)
                _set_cell_text(cells[3], f"{int(n.get('item_done', 0))}/{int(n.get('item_total', 0))}",
                               size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                _set_cell_text(cells[4], str(len(n.get("finding_brief") or [])), size=9,
                               align=WD_ALIGN_PARAGRAPH.CENTER)
        _para(doc, "", space_after=4)

        # S4 汇总
        dedup = self.result.get("dedup") or {}
        review = self.result.get("review") or {}
        s4_lines = [
            f"原始命中 {int(self.summary.get('findings_raw', 0))} 条 → 去重合并 {int(self.summary.get('dedup_merged', 0))} 条"
            f" → 最终 {int(self.summary.get('total', 0))} 条",
        ]
        if review:
            s4_lines.append(
                f"AI 复核：剔除误报 {int(self.summary.get('review_removed', 0))} 条、"
                f"补充漏报 {int(self.summary.get('review_added', 0))} 条")
        _para(doc, "；".join(s4_lines), size=10, line_spacing=1.5, first_line_indent=21)

    # ------------------------------------------------------------------
    def _overall_stats(self, doc: Any) -> None:
        _h1(doc, "三、整体统计汇总")
        sev = {k: int(self.summary.get(k, 0)) for k in ("high", "medium", "low")}
        total = int(self.summary.get("total", 0))
        by_origin = self.summary.get("by_origin") or {}

        _h2(doc, "3.1 问题级别统计")
        table = doc.add_table(rows=1, cols=5)
        table.style = "Table Grid"
        _table_header(table, ["问题总数", "严重", "一般", "轻微", "来源(规则/AI/复核)"],
                      [3.0, 3.0, 3.0, 3.0, 4.0])
        cells = table.add_row().cells
        vals = [str(total), str(sev["high"]), str(sev["medium"]), str(sev["low"]),
                f"{by_origin.get('rule', 0)} / {by_origin.get('ai', 0)} / {by_origin.get('review', 0)}"]
        for idx, val in enumerate(vals):
            color = COLOR_HIGH if idx == 1 and sev["high"] > 0 else None
            _set_cell_text(cells[idx], val, size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=color)
        _para(doc, "", space_after=6)

        by_node = self.summary.get("by_node") or {}
        if by_node:
            _h2(doc, "3.2 各节点问题分布")
            table2 = doc.add_table(rows=1, cols=3)
            table2.style = "Table Grid"
            _table_header(table2, ["序号", "核查节点", "问题数"], [1.8, 10.2, 2.0])
            for i, (name, cnt) in enumerate(sorted(by_node.items(), key=lambda x: -x[1]), start=1):
                cells = table2.add_row().cells
                _set_cell_text(cells[0], str(i), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                _set_cell_text(cells[1], name, size=9)
                _set_cell_text(cells[2], str(cnt), size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
                if i % 2 == 0:
                    for c in cells:
                        _shade_cell(c, SHADE_SUBTLE)
        _para(doc, "", space_after=6)

        concl = (f"本次共执行 {int(self.summary.get('node_done', 0))} 个核查节点、"
                 f"{int(self.summary.get('item_done', 0))} 项核查事项，"
                 f"最终发现问题 {total} 条"
                 f"（严重 {sev['high']} 条、一般 {sev['medium']} 条、轻微 {sev['low']} 条）。")
        _para(doc, concl, size=10.5, line_spacing=1.5, first_line_indent=21)

    # ------------------------------------------------------------------
    def _issue_details(self, doc: Any) -> None:
        doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
        _h1(doc, "四、问题明细清单")
        if not self.findings:
            _para(doc, "本次核查未发现任何问题，文档规范性良好。", size=10.5,
                  line_spacing=1.5, first_line_indent=21)
            return

        # 按节点分组，组内按 严重度→文件→位置 排序
        groups: Dict[str, List[Dict[str, Any]]] = {}
        order: List[str] = []
        for f in self.findings:
            if f.get("state") == "ignored":
                continue
            key = f"{f.get('node_no', '')}{f.get('node_name', '')}"
            if key not in groups:
                groups[key] = []
                order.append(key)
            groups[key].append(f)
        for k in groups:
            groups[k].sort(key=lambda x: (SEVERITY_RANK.get(x.get("severity"), 9),
                                          x.get("file", ""), x.get("location", "")))

        seq = 0
        for gi, key in enumerate(order, start=1):
            items = groups[key]
            _h2(doc, f"4.{gi} {key}（{len(items)} 条）")
            table = doc.add_table(rows=1, cols=7)
            table.style = "Table Grid"
            _table_header(table,
                          ["序号", "级别", "位置", "核查事项", "问题说明", "原文", "整改建议"],
                          [1.0, 1.3, 3.0, 2.4, 3.6, 2.4, 3.3])
            for it in items:
                seq += 1
                sev = it.get("severity", "low")
                cells = table.add_row().cells
                _set_cell_text(cells[0], str(seq), size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
                _set_cell_text(cells[1], SEVERITY_LABEL.get(sev, sev), size=8.5,
                               align=WD_ALIGN_PARAGRAPH.CENTER,
                               color=SEVERITY_COLOR.get(sev, COLOR_TEXT))
                _set_cell_text(cells[2], it.get("location") or "—", size=8.5)
                _set_cell_text(cells[3], it.get("item_title") or "—", size=8.5)
                detail = it.get("detail") or ""
                if int(it.get("merged_count", 1) or 1) > 1:
                    detail += f"\n（已合并 {it.get('merged_count')} 处重复命中）"
                _set_cell_text(cells[4], detail, size=8.5)
                _set_cell_text(cells[5], it.get("snippet") or "—", size=8.5)
                _set_cell_text(cells[6], it.get("suggestion") or "—", size=8.5)
                if it.get("state") == "checked":
                    for c in cells:
                        _shade_cell(c, SHADE_CHECKED)
            _para(doc, "", space_after=8)

    # ------------------------------------------------------------------
    def _dedup_review(self, doc: Any) -> None:
        _h1(doc, "五、去重与 AI 复核说明")
        dedup = self.result.get("dedup") or {}
        review = self.result.get("review") or {}

        _h2(doc, "5.1 跨节点去重")
        _para(doc,
              f"同一问题可能被多个节点（如报告↔说明、说明↔收益法/成本法/市场法联查）同时命中。"
              f"本系统按「同文件 + 同位置 + 同问题性质」自动归并，本次共合并 "
              f"{int(self.summary.get('dedup_merged', 0))} 处重复命中，最终仅保留一条主记录。",
              size=10.5, line_spacing=1.5, first_line_indent=21)

        _h2(doc, "5.2 AI 全局复核")
        if review and review.get("ok"):
            _para(doc,
                  f"AI 复核已启用：剔除误报 {int(self.summary.get('review_removed', 0))} 条、"
                  f"补充漏报 {int(self.summary.get('review_added', 0))} 条。",
                  size=10.5, line_spacing=1.5, first_line_indent=21)
            concl = review.get("conclusion") or ""
            if concl:
                _para(doc, f"复核结论：{concl}", size=10.5, line_spacing=1.5,
                      first_line_indent=21, space_before=4)
        else:
            _para(doc,
                  "本次核查未启用 AI 复核（或 AI 不可用），以下结论仅基于本地确定性规则引擎的命中结果，"
                  "建议对严重与一般级别问题由人工复核确认。",
                  size=10.5, line_spacing=1.5, first_line_indent=21)

    # ------------------------------------------------------------------
    def _suggestions(self, doc: Any) -> None:
        _h1(doc, "六、整改建议与结论")
        _h2(doc, "6.1 总体整改要求")
        general = [
            "按「严重 → 一般 → 轻微」的顺序推进整改，严重问题应在归档前全部清零。",
            "整改应在原始文档副本上进行，保留修改痕迹与版本记录，便于复核与追溯。",
            "含 #REF! 等硬性错误的 Excel 底稿须优先修复公式链，再重新核查。",
            "整改完成后使用本工具重新检测，确认问题闭环后再进入归档或报送流程。",
            "涉密文档整改全过程应在本机离线环境完成，禁止通过网络传输或外发。",
        ]
        for idx, text in enumerate(general, start=1):
            _para(doc, f"{idx}. {text}", size=10.5, line_spacing=1.5, space_after=3)

        _h2(doc, "6.2 结论")
        high = int(self.summary.get("high", 0))
        total = int(self.summary.get("total", 0))
        if self.result.get("status") == "blocked":
            concl = "预处理阶段发现硬性阻断错误（如 #REF!），核查已终止，须先修复源文件后重新核查。"
        elif high > 0:
            concl = ("本次核查发现严重级别问题，相关文档暂不具备归档 / 报送条件，"
                     "请优先完成严重问题整改并重新检测。")
        elif total > 0:
            concl = ("本次核查未发现严重级别问题，但存在一般与轻微问题，"
                     "建议在正式归档 / 报送前完成规范化整理。")
        else:
            concl = "本次核查未发现任何问题，文档规范性良好，具备归档 / 报送条件。"
        _para(doc, concl, size=10.5, bold=True, line_spacing=1.5, first_line_indent=21)

    # ------------------------------------------------------------------
    def _footer_note(self, doc: Any) -> None:
        _para(doc, "", space_after=20)
        _para(doc, "— 报告结束 —", size=10, align=WD_ALIGN_PARAGRAPH.CENTER,
              color=COLOR_MUTED, space_after=16)
        _para(doc, f"报告生成时间：{self.now.strftime('%Y-%m-%d %H:%M:%S')}",
              size=9, align=WD_ALIGN_PARAGRAPH.RIGHT, color=COLOR_MUTED, space_after=2)
        _para(doc, "生成工具：资产评估报告智能核查（本地离线版）",
              size=9, align=WD_ALIGN_PARAGRAPH.RIGHT, color=COLOR_MUTED, space_after=2)
        _para(doc, "保密声明：本报告及检测过程全程离线，未产生任何网络传输记录。",
              size=9, align=WD_ALIGN_PARAGRAPH.RIGHT, color=COLOR_MUTED, space_after=0)

    # ------------------------------------------------------------------
    @staticmethod
    def _status_label(status: Optional[str]) -> str:
        return {
            "done": "已完成", "skipped": "已跳过", "running": "运行中",
            "aborted": "已中断", "failed": "失败", "blocked": "已阻断",
            "canceled": "已取消", "pending": "待执行",
        }.get(status or "", status or "—")


def build_valuation_report(result: Dict[str, Any], out_path: str,
                           checklist: Optional[Dict[str, Any]] = None,
                           operator: str = "", org: str = "",
                           include_cover: bool = True) -> str:
    """函数式入口：由流水线结果生成 Word 终稿报告并返回保存路径。"""
    builder = ValuationReportBuilder(result, checklist, operator, org, include_cover)
    return builder.build(out_path)


def default_report_name() -> str:
    """默认报告文件名（带时间戳，避免覆盖）。"""
    return f"资产评估核查报告_{datetime.now().strftime('%Y%m%d_%H%M%S')}.docx"


if __name__ == "__main__":
    # 自检：用样本跑通流水线并生成报告，验证不崩溃
    import sys
    import tempfile
    import json
    sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
    from checkers.valuation.pipeline import Pipeline  # noqa: E402
    from tests.make_val_samples import EXPECTED  # noqa: E402

    sample_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))),
                              "tests", "samples", "valuation")
    slots = {s: {"name": n, "path": os.path.join(sample_dir, n)}
             for s, n in {"report": "资产评估报告.docx", "statement": "资产评估说明.docx",
                          "income": "收益法测算表.xlsx", "cost": "成本法测算表.xlsx",
                          "market": "市场法测算表.xlsx"}.items()}
    pipe = Pipeline({"slots": slots, "workdir": tempfile.mkdtemp(), "use_ai": False},
                   {"enabled": False}, log=print)
    res = pipe.run()
    checklist = json.load(open(os.path.join(os.path.dirname(os.path.abspath(__file__)),
                                            "..", "..", "config", "valuation", "checklist.json"),
                               encoding="utf-8"))
    out = os.path.join(tempfile.gettempdir(), default_report_name())
    build_valuation_report(res, out, checklist, operator="自检", org="示例单位")
    print("报告已生成：", out, "问题数：", res["summary"].get("total"))
