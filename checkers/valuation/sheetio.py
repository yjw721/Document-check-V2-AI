# -*- coding: utf-8 -*-
"""
Excel 工作表读取层（评估核查专用）
=========================================================================
为预处理（#REF! 检测）与核查规则引擎（勾稽 / 加总 / 跨表同源）提供统一的
工作表访问模型：

    - SheetModel：一张表的「值矩阵 + 公式矩阵 + 合并单元格映射」
    - WorkbookModel：一个工作簿的多表模型

设计要点：
    * 合并单元格在读取时**补齐值**（左上角值填充到整个合并区），解决
      「合并单元格影响解析」的问题，且不修改原始文件；
    * 同时加载「缓存值」与「公式」，使勾稽检查既能取数值、也能检查公式一致性；
    * 全部本地计算，零联网。
"""

from __future__ import annotations

from typing import Any, Dict, List, Optional, Tuple

try:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter, range_boundaries
except Exception:  # pragma: no cover - 环境缺依赖时由调用方降级
    load_workbook = None  # type: ignore
    get_column_letter = None  # type: ignore
    range_boundaries = None  # type: ignore


# ---------------------------------------------------------------------------
# 数值 / 文本工具
# ---------------------------------------------------------------------------
def to_num(v: Any) -> Optional[float]:
    """把单元格值转为 float；非数值返回 None。"""
    if v is None:
        return None
    if isinstance(v, bool):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).strip()
    if not s:
        return None
    s = s.replace(",", "").replace("，", "")
    neg = False
    if s.startswith("(") and s.endswith(")"):
        neg, s = True, s[1:-1]
    if s.endswith("%"):
        try:
            val = float(s[:-1])
            return -val / 100.0 if neg else val / 100.0
        except ValueError:
            return None
    s = s.replace("¥", "").replace("￥", "").replace("$", "")
    try:
        val = float(s)
    except ValueError:
        return None
    return -val if neg else val


def is_num(v: Any) -> bool:
    return to_num(v) is not None


def is_error_value(v: Any) -> bool:
    """判断是否为 Excel 错误值（#REF! / #VALUE! / #DIV/0! 等）。"""
    if not isinstance(v, str):
        return False
    s = v.strip()
    return s.upper() in ("#REF!", "#VALUE!", "#DIV/0!", "#NAME?", "#N/A",
                         "#NULL!", "#NUM!", "#SPILL!", "#CALC!", "#GETTING_DATA")


def fmt_num(v: Optional[float]) -> str:
    if v is None:
        return "-"
    if abs(v) >= 1e8 or (abs(v) < 1e-4 and v != 0):
        return f"{v:.4g}"
    return f"{v:,.2f}".rstrip("0").rstrip(".")


def cell_addr(row: int, col: int) -> str:
    """(1-based row, col) → A1 地址。"""
    if get_column_letter is None:
        return f"R{row}C{col}"
    return f"{get_column_letter(col)}{row}"


# ---------------------------------------------------------------------------
# 工作表模型
# ---------------------------------------------------------------------------
class SheetModel:
    """一张工作表的轻量模型（值矩阵 + 公式矩阵 + 合并映射）。"""

    def __init__(self, name: str):
        self.name = name
        self.values: List[List[Any]] = []      # 1-based 访问：values[r-1][c-1]
        self.formulas: List[List[Optional[str]]] = []
        self.nrows: int = 0
        self.ncols: int = 0
        self.print_area: str = ""
        self.state: str = "visible"            # visible / hidden / veryHidden
        self.merged: List[str] = []

    # ---------- 基础访问 ----------
    def val(self, row: int, col: int) -> Any:
        if 1 <= row <= self.nrows and 1 <= col <= self.ncols:
            return self.values[row - 1][col - 1]
        return None

    def num(self, row: int, col: int) -> Optional[float]:
        return to_num(self.val(row, col))

    def formula(self, row: int, col: int) -> Optional[str]:
        if 1 <= row <= self.nrows and 1 <= col <= self.ncols:
            return self.formulas[row - 1][col - 1]
        return None

    def row_values(self, row: int) -> List[Any]:
        if 1 <= row <= self.nrows:
            return list(self.values[row - 1])
        return []

    def col_values(self, col: int) -> List[Any]:
        if 1 <= col <= self.ncols:
            return [self.values[r - 1][col - 1] for r in range(1, self.nrows + 1)]
        return []

    def row_nums(self, row: int) -> List[Tuple[int, float]]:
        """返回该行的 (列号, 数值) 列表。"""
        out = []
        for c, v in enumerate(self.row_values(row), start=1):
            n = to_num(v)
            if n is not None:
                out.append((c, n))
        return out

    def text(self, max_rows: int = 300, max_cols: int = 40) -> str:
        """导出为文本，供 AI 核查使用（行号 + 内容）。"""
        lines: List[str] = []
        for r in range(1, min(self.nrows, max_rows) + 1):
            cells = []
            for c in range(1, min(self.ncols, max_cols) + 1):
                v = self.val(r, c)
                if v is None:
                    continue
                s = str(v).replace("\n", " ").strip()
                if s:
                    cells.append(f"{cell_addr(r, c)}={s}")
            if cells:
                lines.append(f"[{r}] " + " | ".join(cells))
        return "\n".join(lines)

    # ---------- 结构识别 ----------
    def find_rows(self, keywords, search_cols: int = 4, limit: int = 0) -> List[int]:
        """在前 search_cols 列中查找包含任一关键词的行号。"""
        hits: List[int] = []
        for r in range(1, self.nrows + 1):
            for c in range(1, min(self.ncols, search_cols) + 1):
                v = self.val(r, c)
                if v is None:
                    continue
                s = str(v)
                if any(k in s for k in keywords):
                    hits.append(r)
                    break
            if limit and len(hits) >= limit:
                break
        return hits

    def find_cols(self, keywords, search_rows: int = 5, limit: int = 0) -> List[int]:
        """在前 search_rows 行中查找包含任一关键词的列号。"""
        hits: List[int] = []
        for c in range(1, self.ncols + 1):
            for r in range(1, min(self.nrows, search_rows) + 1):
                v = self.val(r, c)
                if v is None:
                    continue
                s = str(v)
                if any(k in s for k in keywords):
                    hits.append(c)
                    break
            if limit and len(hits) >= limit:
                break
        return hits

    def to_dict(self) -> Dict[str, Any]:
        return {
            "name": self.name,
            "nrows": self.nrows,
            "ncols": self.ncols,
            "state": self.state,
            "print_area": self.print_area,
            "merged": len(self.merged),
        }


class WorkbookModel:
    """一个工作簿的多表模型。"""

    def __init__(self, path: str):
        self.path = path
        self.sheets: Dict[str, SheetModel] = {}
        self.order: List[str] = []
        self.load_error: str = ""

    def get(self, name: str) -> Optional[SheetModel]:
        return self.sheets.get(name)

    def sheet_names(self) -> List[str]:
        return list(self.order)

    def text(self, sheets: Optional[List[str]] = None,
             max_rows: int = 300, max_cols: int = 40,
             max_chars: int = 0) -> str:
        """导出指定表（默认全部）的文本，供 AI 核查使用。"""
        names = sheets or self.order
        parts: List[str] = []
        for name in names:
            sm = self.sheets.get(name)
            if not sm:
                continue
            body = sm.text(max_rows, max_cols)
            if body:
                parts.append(f"### Sheet：{name}\n{body}")
        out = "\n\n".join(parts)
        if max_chars and len(out) > max_chars:
            out = out[:max_chars] + "\n…（内容过长已截断）"
        return out


# ---------------------------------------------------------------------------
# 加载
# ---------------------------------------------------------------------------
_MAX_CELLS_PER_SHEET = 400000       # 单表扫描上限，防止超大表拖垮进程
_ERR_CACHE: Dict[str, Any] = {}


def load_workbook_model(path: str, with_formula: bool = True) -> WorkbookModel:
    """加载工作簿模型（值 + 公式）。失败时 load_error 非空。"""
    model = WorkbookModel(path)
    if load_workbook is None:
        model.load_error = "openpyxl 组件缺失，无法解析 Excel"
        return model

    try:
        wb_v = load_workbook(path, data_only=True, read_only=False)
    except Exception as exc:  # noqa: BLE001
        model.load_error = f"工作簿解析失败：{exc}"
        return model

    wb_f = None
    if with_formula:
        try:
            wb_f = load_workbook(path, data_only=False, read_only=False)
        except Exception:  # noqa: BLE001 - 公式矩阵加载失败不影响值核查
            wb_f = None

    try:
        for ws in wb_v.worksheets:
            sm = SheetModel(ws.title)
            sm.state = str(getattr(ws, "sheet_state", "visible") or "visible")
            sm.print_area = str(getattr(ws, "print_area", "") or "")
            try:
                sm.merged = [str(r) for r in ws.merged_cells.ranges]
            except Exception:  # noqa: BLE001
                sm.merged = []

            nrows = int(getattr(ws, "max_row", 0) or 0)
            ncols = int(getattr(ws, "max_column", 0) or 0)
            # 合并单元格补齐：把左上角值填充到整个区域
            fill: Dict[Tuple[int, int], Any] = {}
            if range_boundaries is not None:
                for rng in ws.merged_cells.ranges:
                    try:
                        c1, r1, c2, r2 = range_boundaries(str(rng))
                        base = ws.cell(row=r1, column=c1).value
                        for rr in range(r1, r2 + 1):
                            for cc in range(c1, c2 + 1):
                                fill[(rr, cc)] = base
                    except Exception:  # noqa: BLE001
                        continue

            values: List[List[Any]] = []
            formulas: List[List[Optional[str]]] = []
            ws_f = None
            if wb_f is not None:
                try:
                    ws_f = wb_f[ws.title]
                except Exception:  # noqa: BLE001
                    ws_f = None

            cap_r = min(nrows, max(1, _MAX_CELLS_PER_SHEET // max(1, ncols)))
            for r in range(1, cap_r + 1):
                vrow: List[Any] = []
                frow: List[Optional[str]] = []
                for c in range(1, ncols + 1):
                    v = ws.cell(row=r, column=c).value
                    if v is None and (r, c) in fill:
                        v = fill[(r, c)]
                    vrow.append(v)
                    f = None
                    if ws_f is not None:
                        try:
                            cv = ws_f.cell(row=r, column=c)
                            if cv.data_type == "f":
                                f = str(cv.value)
                        except Exception:  # noqa: BLE001
                            f = None
                    frow.append(f)
                values.append(vrow)
                formulas.append(frow)

            sm.values = values
            sm.formulas = formulas
            sm.nrows = len(values)
            sm.ncols = ncols
            model.sheets[ws.title] = sm
            model.order.append(ws.title)
    finally:
        try:
            wb_v.close()
        except Exception:  # noqa: BLE001
            pass
        if wb_f is not None:
            try:
                wb_f.close()
            except Exception:  # noqa: BLE001
                pass
    return model


# ---------------------------------------------------------------------------
# 打印区域
# ---------------------------------------------------------------------------
def print_area_ranges(ws: Any, sm: SheetModel) -> List[Tuple[int, int, int, int]]:
    """解析打印区域为 (min_row, min_col, max_row, max_col)；未设置时返回整表已用区。"""
    area = str(getattr(ws, "print_area", "") or "")
    out: List[Tuple[int, int, int, int]] = []
    if area and range_boundaries is not None:
        for seg in area.split(","):
            seg = seg.strip()
            if not seg:
                continue
            try:
                c1, r1, c2, r2 = range_boundaries(seg.replace("$", ""))
                out.append((r1, c1, r2, c2))
            except Exception:  # noqa: BLE001
                continue
    if not out:
        if sm.nrows and sm.ncols:
            out.append((1, 1, sm.nrows, sm.ncols))
    return out
