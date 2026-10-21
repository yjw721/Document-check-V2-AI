# -*- coding: utf-8 -*-
"""
阶段一 · 文件上传与预处理
=========================================================================
对应需求文档「阶段一 · 预处理详解」：

    1. 上传 5 个附件（评估报告 / 评估说明 + 三法测算表）
    2. 后台脚本清理 Excel（合并单元格补齐、格式残留归一化）
    3. #REF! 错误检查（硬性闸门）
       —— 打印区域内存在 #REF! 即停止报错，交回用户修复
    4. 深度清理：删除隐藏表 / 清理已定义名称 / 断开外部链接

设计原则（与本项目一致）：
    * 只读取、只产出「清理后的副本」，绝不修改用户原始文件；
    * 不需要清理时直接复制原文件，零改写风险；
    * 全程本地计算，零联网。
"""

from __future__ import annotations

import os
import re
import shutil
import zipfile
from typing import Any, Dict, List, Optional, Tuple

from .sheetio import load_workbook_model, is_error_value, cell_addr, to_num

try:
    from openpyxl import load_workbook
    from openpyxl.utils import range_boundaries
except Exception:  # pragma: no cover
    load_workbook = None  # type: ignore
    range_boundaries = None  # type: ignore


# 硬性拦截的错误值
BLOCK_ERROR = "#REF!"
# 其它错误值：记录但不拦截
OTHER_ERRORS = ("#VALUE!", "#DIV/0!", "#NAME?", "#N/A", "#NULL!", "#NUM!",
                "#SPILL!", "#CALC!")


# ---------------------------------------------------------------------------
# 工具
# ---------------------------------------------------------------------------
def _parse_area(area: str, nrows: int, ncols: int) -> List[Tuple[int, int, int, int]]:
    """解析打印区域字符串为 (min_row, min_col, max_row, max_col)。"""
    out: List[Tuple[int, int, int, int]] = []
    area = (area or "").strip()
    if area and range_boundaries is not None:
        for seg in area.split(","):
            seg = seg.strip().replace("$", "")
            if not seg:
                continue
            try:
                c1, r1, c2, r2 = range_boundaries(seg)
                out.append((r1, c1, r2, c2))
            except Exception:  # noqa: BLE001
                continue
    if not out and nrows and ncols:
        out.append((1, 1, nrows, ncols))
    # 夹紧到实际范围，避免越界
    clipped = []
    for r1, c1, r2, c2 in out:
        clipped.append((max(1, r1), max(1, c1), min(nrows, r2), min(ncols, c2)))
    return clipped


def _scan_errors(model, only_block: bool = False) -> List[Dict[str, Any]]:
    """扫描工作簿中的错误单元格（限定打印区域内）。

    同时检查：
      - 公式文本中的 #REF!（公式链损坏）
      - 缓存值中的错误值（#REF! / #VALUE! …）
    """
    hits: List[Dict[str, Any]] = []
    for name in model.order:
        sm = model.sheets[name]
        ranges = _parse_area(sm.print_area, sm.nrows, sm.ncols)
        if not ranges:
            continue
        seen = set()
        for r1, c1, r2, c2 in ranges:
            for r in range(r1, r2 + 1):
                for c in range(c1, c2 + 1):
                    val = sm.val(r, c)
                    fml = sm.formula(r, c)
                    err = ""
                    if isinstance(val, str) and is_error_value(val):
                        err = val.strip().upper()
                    elif fml and BLOCK_ERROR in fml.upper():
                        err = BLOCK_ERROR
                    if not err:
                        continue
                    if only_block and err != BLOCK_ERROR:
                        continue
                    key = (r, c)
                    if key in seen:
                        continue
                    seen.add(key)
                    hits.append({
                        "sheet": name,
                        "cell": cell_addr(r, c),
                        "row": r, "col": c,
                        "error": err,
                        "formula": (fml or "")[:200],
                        "in_print_area": bool(sm.print_area),
                    })
                    if len(hits) >= 500:      # 上限保护
                        return hits
    return hits


# ---------------------------------------------------------------------------
# 深度清理
# ---------------------------------------------------------------------------
def _strip_external_links(src: str, dst: str) -> int:
    """在 ZIP 层移除 externalLinks 部件及其关系引用，返回移除数量。

    仅重写必要的 XML 部件，不触碰工作表数据。
    """
    removed = 0
    try:
        with zipfile.ZipFile(src, "r") as zin:
            names = zin.namelist()
            keep = [n for n in names if not n.startswith("xl/externalLinks/")]
            removed = len(names) - len(keep)
            if removed == 0:
                return 0
            payload: Dict[str, bytes] = {}
            for n in keep:
                data = zin.read(n)
                if n == "xl/_rels/workbook.xml.rels":
                    data = re.sub(
                        rb"<Relationship\b[^>]*externalLink[^>]*/>", b"", data)
                elif n == "[Content_Types].xml":
                    data = re.sub(
                        rb"<Override\b[^>]*externalLinks/[^>]*/>", b"", data)
                payload[n] = data
        with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
            for n, data in payload.items():
                zout.writestr(n, data)
        return removed
    except Exception:  # noqa: BLE001 - 清理失败不阻断流程
        try:
            shutil.copyfile(src, dst)
        except Exception:  # noqa: BLE001
            pass
        return 0


def deep_clean_excel(src: str, dst: str) -> Dict[str, Any]:
    """深度清理：删除隐藏表 / 清理已定义名称 / 断开外部链接。

    返回清理统计；无需清理时直接复制原文件（零改写风险）。
    """
    stat: Dict[str, Any] = {
        "hidden_sheets": [], "defined_names": 0, "external_links": 0,
        "error": "", "rewritten": False,
    }
    if load_workbook is None:
        shutil.copyfile(src, dst)
        stat["error"] = "openpyxl 缺失，已跳过深度清理"
        return stat

    tmp = dst + ".tmp.xlsx"
    try:
        wb = load_workbook(src)
    except Exception as exc:  # noqa: BLE001
        shutil.copyfile(src, dst)
        stat["error"] = f"工作簿打开失败，已跳过深度清理：{exc}"
        return stat

    changed = False
    # 1) 删除隐藏工作表（至少保留一张表）
    try:
        visible = [ws for ws in wb.worksheets
                   if str(getattr(ws, "sheet_state", "visible")) == "visible"]
        if visible and len(visible) != len(wb.worksheets):
            hidden = [ws.title for ws in wb.worksheets
                      if str(getattr(ws, "sheet_state", "visible")) != "visible"]
            for title in hidden:
                try:
                    del wb[title]
                    stat["hidden_sheets"].append(title)
                    changed = True
                except Exception:  # noqa: BLE001
                    continue
    except Exception:  # noqa: BLE001
        pass

    # 2) 清理已定义名称
    try:
        dn = getattr(wb, "defined_names", None)
        if dn is not None:
            names = list(dn.keys()) if hasattr(dn, "keys") else list(dn)
            for nm in names:
                try:
                    del dn[nm]
                    stat["defined_names"] += 1
                    changed = True
                except Exception:  # noqa: BLE001
                    try:
                        dn.pop(nm, None)
                        stat["defined_names"] += 1
                        changed = True
                    except Exception:  # noqa: BLE001
                        continue
    except Exception:  # noqa: BLE001
        pass

    if changed:
        try:
            wb.save(tmp)
            stat["rewritten"] = True
        except Exception as exc:  # noqa: BLE001
            stat["error"] = f"清理后保存失败，已回退原文件：{exc}"
            shutil.copyfile(src, tmp)
    else:
        shutil.copyfile(src, tmp)

    # 3) 断开外部链接（ZIP 层）
    try:
        with zipfile.ZipFile(tmp, "r") as zin:
            ext = [n for n in zin.namelist() if n.startswith("xl/externalLinks/")]
        stat["external_links"] = len(ext)
        if ext:
            _strip_external_links(tmp, dst)
        else:
            shutil.copyfile(tmp, dst)
    except Exception as exc:  # noqa: BLE001
        stat["error"] = (stat["error"] + f"；外部链接检查失败：{exc}").strip("；")
        shutil.copyfile(tmp, dst)

    if os.path.exists(tmp):
        try:
            os.remove(tmp)
        except OSError:
            pass
    return stat


# ---------------------------------------------------------------------------
# 单文件预处理
# ---------------------------------------------------------------------------
def preprocess_excel(src_path: str, out_path: str, file_name: str) -> Dict[str, Any]:
    """预处理单个 Excel：脚本清理 → #REF! 硬拦截 → 深度清理。"""
    res: Dict[str, Any] = {
        "file": file_name, "kind": "excel", "status": "ok",
        "blocked": False, "message": "",
        "ref_errors": [], "other_errors": [], "clean": {},
        "sheets": [], "out_path": "",
    }

    model = load_workbook_model(src_path)
    if model.load_error:
        res["status"] = "error"
        res["message"] = model.load_error
        return res

    res["sheets"] = [model.sheets[n].to_dict() for n in model.order]

    # 步骤 2：脚本清理（合并单元格 / 格式残留）——统计，不改写原文件
    merged = sum(len(model.sheets[n].merged) for n in model.order)
    res["merged_ranges"] = merged

    # 步骤 3：#REF! 硬性闸门
    block_hits = _scan_errors(model, only_block=True)
    other_hits = [h for h in _scan_errors(model) if h["error"] != BLOCK_ERROR]
    res["ref_errors"] = block_hits[:50]
    res["other_errors"] = other_hits[:50]

    if block_hits:
        res["blocked"] = True
        res["status"] = "blocked"
        first = block_hits[0]
        res["message"] = (
            f"打印区域内检测到 {len(block_hits)} 处 #REF! 引用错误"
            f"（首处：{first['sheet']}!{first['cell']}）。"
            "公式链已损坏、测算数据不可信，已按规则停止核查，请修复后重新上传。"
        )
        return res

    # 步骤 4：深度清理
    clean = deep_clean_excel(src_path, out_path)
    res["clean"] = clean
    res["out_path"] = out_path
    if other_hits:
        res["message"] = f"检测到 {len(other_hits)} 处非 #REF! 错误值（如 #VALUE!/#DIV/0!），已记录，不阻断核查"
    else:
        res["message"] = "预处理通过：打印区域无 #REF! 错误"
    if clean.get("error"):
        res["message"] = (res["message"] + "；" + clean["error"]).strip("；")
    return res


def preprocess_doc(src_path: str, out_path: str, file_name: str) -> Dict[str, Any]:
    """预处理文档类附件（评估报告 / 评估说明）：只做可读性与规模校验。"""
    res: Dict[str, Any] = {
        "file": file_name, "kind": "doc", "status": "ok", "blocked": False,
        "message": "", "paragraphs": 0, "tables": 0, "chars": 0, "out_path": out_path,
    }
    try:
        shutil.copyfile(src_path, out_path)
    except Exception as exc:  # noqa: BLE001
        res["status"] = "error"
        res["message"] = f"文件复制失败：{exc}"
        return res
    try:
        import docx  # python-docx
        d = docx.Document(out_path)
        res["paragraphs"] = len(d.paragraphs)
        res["tables"] = len(d.tables)
        res["chars"] = sum(len(p.text or "") for p in d.paragraphs)
        res["message"] = f"解析成功：{res['paragraphs']} 段 / {res['tables']} 张表 / {res['chars']} 字"
    except Exception as exc:  # noqa: BLE001
        res["status"] = "error"
        res["message"] = f"文档解析失败：{exc}"
    return res


def run_preprocess(slot_files: Dict[str, Dict[str, str]],
                   cache_dir: str) -> Dict[str, Any]:
    """对全部槽位执行阶段一预处理。

    slot_files: {slot: {"name": 原文件名, "path": 上传落盘路径}}
    返回 {items: [...], blocked: bool, message: str}
    """
    os.makedirs(cache_dir, exist_ok=True)
    items: List[Dict[str, Any]] = []
    blocked = False
    blocked_files: List[str] = []

    for slot, info in (slot_files or {}).items():
        src = info.get("path") or ""
        name = info.get("name") or os.path.basename(src)
        if not src or not os.path.exists(src):
            continue
        out = os.path.join(cache_dir, f"{slot}__{name}")
        if slot in ("income", "cost", "market"):
            item = preprocess_excel(src, out, name)
        else:
            item = preprocess_doc(src, out, name)
        item["slot"] = slot
        items.append(item)
        if item.get("blocked"):
            blocked = True
            blocked_files.append(name)

    message = ""
    if blocked:
        message = (
            "⛔ 阶段一硬拦截：以下测算表打印区域内存在 #REF! 错误，核查已停止——"
            + "、".join(blocked_files)
            + "。请修复公式后重新上传。"
        )
    else:
        message = "阶段一预处理完成：全部附件通过 #REF! 检查，已完成深度清理（删隐藏表 / 清名称 / 断外链）"
    return {"items": items, "blocked": blocked, "message": message}
