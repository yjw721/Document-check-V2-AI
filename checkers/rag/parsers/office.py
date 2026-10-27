# -*- coding: utf-8 -*-
"""Office / 纯文本解析。

- ``.txt / .md``  → 逐行块
- ``.csv``        → 逐行块（列用 `` | `` 连接，保留表头语义）
- ``.docx``       → 段落块 + 表格行块
- ``.xlsx``       → 「工作表名 + 行号」块
- ``.doc``（旧版二进制）→ 尽力按文本抽取，失败给出可操作的提示
"""

from __future__ import annotations

import io
import re
from typing import List

from checkers.rag.parsers.base import Block, decode_text, normalize

_CSV_SPLIT = re.compile(r",(?=(?:[^\"]*\"[^\"]*\")*[^\"]*$)")


def parse_plain(name: str, raw: bytes) -> List[Block]:
    """.txt / .md 解析。"""
    text = normalize(decode_text(raw))
    out: List[Block] = []
    for i, ln in enumerate(text.split("\n"), 1):
        s = ln.strip()
        if len(s) >= 2:
            out.append((f"第 {i} 行", s))
    return out


def parse_csv(name: str, raw: bytes) -> List[Block]:
    """.csv 解析：保留列结构，去掉纯分隔行。"""
    text = normalize(decode_text(raw))
    out: List[Block] = []
    for i, ln in enumerate(text.split("\n"), 1):
        s = ln.strip()
        if len(s) < 2:
            continue
        if re.fullmatch(r"[-=+|\s]+", s):
            continue
        cells = [c.strip().strip('"') for c in _CSV_SPLIT.split(s)]
        cells = [c for c in cells if c]
        if not cells:
            continue
        out.append((f"第 {i} 行", " | ".join(cells)))
    return out


def parse_docx(name: str, raw: bytes) -> List[Block]:
    """.docx 解析：按文档流顺序混合段落与表格行。

    旧实现先遍历 ``doc.paragraphs``、再遍历 ``doc.tables``，于是所有表格被
    整体挪到文末，脱离其所属标题与引导句——章节路径 / 上下文前缀随之错位
    （如 LPR 表格被挂在「三、定期存款」标题下，按标题检索全部失准）。
    这里改为按 ``w:body`` 的真实子元素顺序遍历，段落与表格交错输出；
    同时对横向合并单元格（``row.cells`` 展开为相邻重复值）做相邻去重。
    """
    from docx import Document  # type: ignore
    from docx.oxml.ns import qn  # type: ignore
    from docx.table import Table  # type: ignore
    from docx.text.paragraph import Paragraph  # type: ignore

    doc = Document(io.BytesIO(raw))
    blocks: List[Block] = []
    pi = ti = 0
    for child in doc.element.body.iterchildren():
        if child.tag == qn("w:p"):
            pi += 1
            t = Paragraph(child, doc).text.strip()
            if t:
                blocks.append((f"第 {pi} 段", t))
        elif child.tag == qn("w:tbl"):
            ti += 1
            table = Table(child, doc)
            for ri, row in enumerate(table.rows, 1):
                cells: List[str] = []
                for c in row.cells:
                    v = c.text.strip()
                    # 横向合并的单元格在 row.cells 中被展开为重复值 → 相邻去重
                    if v and (not cells or v != cells[-1]):
                        cells.append(v)
                if cells:
                    blocks.append((f"表 {ti} 第 {ri} 行", " | ".join(cells)))
    return blocks


def parse_xlsx(name: str, raw: bytes) -> List[Block]:
    """.xlsx 解析：逐行，位置标识含工作表名。"""
    import openpyxl  # type: ignore

    wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    blocks: List[Block] = []
    try:
        for ws in wb.worksheets:
            for ri, row in enumerate(ws.iter_rows(values_only=True), 1):
                cells = [str(v).strip() for v in row if v is not None and str(v).strip()]
                if not cells:
                    continue
                t = " | ".join(cells)
                if len(t) >= 2:
                    blocks.append((f"「{ws.title}」第 {ri} 行", t))
    finally:
        wb.close()
    return blocks


def parse_doc(name: str, raw: bytes) -> List[Block]:
    """.doc（旧版二进制）兜底抽取。

    python-docx 不支持 .doc。这里做两件事：
    1. 用 antiword / catdoc 之类的外部工具（若存在）；
    2. 否则在二进制流里按 UTF-16LE / GBK 提取可读中文片段。
    仍失败则抛出带可执行建议的异常，界面上会直接提示转存为 .docx。
    """
    import subprocess  # noqa: S404 - 仅在用户本机调用外部转换器

    for cmd in ("antiword", "catdoc"):
        try:
            out = subprocess.run([cmd, "-"], input=raw, capture_output=True, timeout=20)
            if out.returncode == 0 and out.stdout:
                text = normalize(out.stdout.decode("utf-8", errors="replace"))
                if text:
                    return [(f"第 {i} 行", ln) for i, ln in enumerate(text.split("\n"), 1)
                            if len(ln.strip()) >= 2]
        except (OSError, subprocess.SubprocessError):
            continue

    blobs: List[str] = []
    for enc in ("utf-16-le", "gb18030"):
        try:
            s = raw.decode(enc, errors="ignore")
        except (UnicodeDecodeError, LookupError):
            continue
        parts = re.findall(r"[\u4e00-\u9fffA-Za-z0-9，。；：、（）%．.\-]{6,}", s)
        if parts:
            blobs.append(" ".join(parts))
    text = normalize(max(blobs, key=len) if blobs else "")
    if not text:
        raise ValueError(
            "无法解析 .doc（旧版二进制格式）。请在 Word 中「另存为 .docx」后重新上传；"
            "或在本机安装 antiword / catdoc 以支持自动转换。"
        )
    return [(f"第 {i} 段", ln) for i, ln in enumerate(text.split("\n"), 1)
            if len(ln.strip()) >= 2]
